import hashlib
import json
from types import SimpleNamespace

import numpy as np
import openpyxl
import pandas as pd
import pytest

from downstream.AffibodyMHC import build_libb_wetlab_candidate_handoff as handoff
from downstream.AffibodyMHC.build_libb_wetlab_candidate_handoff import (
    DEFAULT_PROPOSED_TARGETS,
    EXPECTED_TARGET_SEQUENCES,
    _forbidden_outcome_columns,
    _validate_target_and_rank_integrity,
    hamming,
    parse_selected_peptides,
    within_target_min_hamming,
)


def test_hamming_and_target_local_minimum():
    assert hamming("AAAAA", "AAAAD") == 1
    assert hamming("AAAAA", "DDDDD") == 5
    frame = pd.DataFrame(
        {
            "peptide_design_code": ["AF", "AF", "AF", "AH"],
            "affibody_design_code": ["AAAAA", "AAAAD", "DDDDD", "AAAAA"],
        },
        index=[4, 2, 9, 1],
    )
    observed = within_target_min_hamming(frame)
    assert observed.to_dict() == {4: 1, 2: 1, 9: 4, 1: 5}


def test_default_batch_is_explicit_neutral_code_coverage_set():
    assert parse_selected_peptides(None) == DEFAULT_PROPOSED_TARGETS
    assert DEFAULT_PROPOSED_TARGETS == (
        "AF", "AH", "DP", "EA", "LV", "MW", "NF", "PH", "TL", "VV"
    )
    assert parse_selected_peptides("VV,TL,PH,NF,MW,LV,EA,DP,AH,AF") == (
        "VV", "TL", "PH", "NF", "MW", "LV", "EA", "DP", "AH", "AF"
    )
    with pytest.raises(ValueError, match="exactly ten"):
        parse_selected_peptides("AF,AH")
    with pytest.raises(ValueError, match="duplicates"):
        parse_selected_peptides("AF,AF,DP,EA,LV,MW,NF,PH,TL,VV")


def test_outcome_column_guard_matches_unlisted_retention_spelling():
    assert _forbidden_outcome_columns(["pair_uid", "observed_retention_percent"]) == [
        "observed_retention_percent"
    ]


def test_exact_target_and_rank_integrity():
    rows = []
    for peptide, sequence in EXPECTED_TARGET_SEQUENCES.items():
        for rank in (1, 2):
            rows.append(
                {
                    "pair_uid": f"{peptide}-{rank}",
                    "peptide_design_code": peptide,
                    "peptide_full_sequence": sequence,
                    "wetlab_rank": rank,
                }
            )
    frame = pd.DataFrame(rows)
    _validate_target_and_rank_integrity(frame, "test")
    with pytest.raises(ValueError, match="exact corrected 12-target"):
        _validate_target_and_rank_integrity(frame.loc[frame.peptide_design_code.ne("VV")], "test")
    bad = frame.copy()
    bad.loc[bad.pair_uid.eq("AF-2"), "wetlab_rank"] = 3
    with pytest.raises(ValueError, match="exactly 1..N"):
        _validate_target_and_rank_integrity(bad, "test")


def _sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _affibody_sequence(code):
    sequence = list(
        "VDNKFAKEFANAAAEIAHLPNLNEEQFDAFVQSLFDDPSQSANLLAEAKKLNDAQAPK"
    )
    for amino_acid, position in zip(code, (6, 10, 13, 14, 17)):
        sequence[position - 1] = amino_acid
    return "".join(sequence)


def _shortlist_rows(prefix, code, score):
    rows = []
    rank_amino_acids = "ADEFIKLMNS"
    mint_extra_first_char = dict(zip(DEFAULT_PROPOSED_TARGETS[:6], "ADFIKL"))
    for peptide, sequence in EXPECTED_TARGET_SEQUENCES.items():
        for rank in range(1, 11):
            ranked_code = code[:4] + rank_amino_acids[rank - 1]
            if prefix == "mint" and rank == 10 and peptide in mint_extra_first_char:
                ranked_code = mint_extra_first_char[peptide] + code[1:4] + ranked_code[-1]
            rows.append(
                {
                    "pair_uid": f"{prefix}-{peptide}-{rank}",
                    "peptide_design_code": peptide,
                    "peptide_full_sequence": sequence,
                    "affibody_design_code": ranked_code,
                    "chain1_smart_hla_linker_peptide_sequence": f"ASSAY-{sequence}",
                    "chain2_affibody_sequence": _affibody_sequence(ranked_code),
                    "wetlab_rank": rank,
                    "ensemble_score": score - rank / 1000,
                    "component_top10_votes": 4,
                    "within_peptide_additive_7site_rank": rank,
                    "within_peptide_mint_layer5_rank": rank,
                    "within_peptide_stab_designed_ordered_rank": rank,
                    "within_peptide_rde_network_designed_3fold_rank": rank,
                    "additive_7site": 0.0,
                    "mint_layer5": 0.0,
                    "stab_designed_ordered": 0.0,
                    "rde_network_designed_3fold": 0.0,
                    "normalized_component_disagreement": 0.0,
                    "observed_in_any_raw_round": False,
                    "observed_in_r009_or_r010": False,
                    "affibody_identity_seen_in_strict_training": False,
                    "high_confidence_weak_negative": False,
                    "prior_weak_label": np.nan,
                    "pooled_r009_r010_count": 0,
                    "r009_count": 0,
                    "r010_count": 0,
                }
            )
    return pd.DataFrame(rows)


def _make_handoff_fixture(tmp_path, monkeypatch):
    monkeypatch.setattr(handoff, "REPO_ROOT", tmp_path)
    private = tmp_path / "private_data"
    private.mkdir()
    universe = private / "universe"
    universe.mkdir()

    comparator_rows = []
    for peptide, sequence in EXPECTED_TARGET_SEQUENCES.items():
        comparator_rows.append(
            {
                "pair_uid": f"comparator-{peptide}",
                "peptide_design_code": peptide,
                "peptide_full_sequence": sequence,
                "affibody_design_code": "DDDDD",
                "chain2_affibody_sequence": _affibody_sequence("DDDDD"),
                "r009_count": 600,
                "r010_count": 400,
                "pooled_r009_r010_count": 1000,
                "excluded_directly_measured": False,
                "excluded_pooled_r009_r010_top2pct": True,
            }
        )
    for index in range(7_536):
        source = comparator_rows[index % len(EXPECTED_TARGET_SEQUENCES)]
        comparator_rows.append(
            {
                **source,
                "pair_uid": f"extra-{index}",
                "r009_count": 1,
                "r010_count": 0,
                "pooled_r009_r010_count": 1,
            }
        )
    excluded_path = universe / "excluded_pairs.parquet"
    pd.DataFrame(comparator_rows).to_parquet(excluded_path, index=False)
    candidate_manifest = {
        "schema_version": "libb-existing-target-selection-missed-candidates-v1",
        "scope": {
            "target_peptide_codes": list(EXPECTED_TARGET_SEQUENCES),
            "target_peptide_full_sequences": EXPECTED_TARGET_SEQUENCES,
        },
        "outputs": {
            "excluded_pairs_rows": len(comparator_rows),
            "excluded_pairs_sha256": _sha256(excluded_path),
        },
        "exclusions": {"pooled_positive_nonmeasured_pairs": len(comparator_rows)},
    }
    candidate_manifest_path = universe / "manifest.json"
    candidate_manifest_path.write_text(json.dumps(candidate_manifest))
    candidate_manifest_sha = _sha256(candidate_manifest_path)

    shortlists = {
        "ensemble": _shortlist_rows("all4", "AAAAA", 0.5),
        "MINT": _shortlist_rows("mint", "EEEEE", 0.6),
        "MINT+StaB": _shortlist_rows("mint-stab", "FFFFF", 0.55),
    }
    for label, lock_id in handoff.EXPECTED_LOCK_IDS.items():
        directory = private / label.replace("+", "_").lower()
        directory.mkdir()
        shortlists[label].to_csv(directory / "wetlab_shortlist.csv", index=False)
        manifest = {
            "schema_version": "libb-ensemble-wetlab-shortlist-v1",
            "ensemble_lock": {
                "lock_id": lock_id,
                "expected_lock_id": lock_id,
                "bundle_manifest_sha256": "b" * 64,
                "bundle_schema_version": "libb-weak-oof-deployment-locks-v2",
            },
            "candidate_universe": {"manifest_sha256": candidate_manifest_sha},
            "outcome_access": {
                "retention_table_read": False,
                "retention_used_for_lock_or_selection": False,
                "retention_columns_allowed_in_component_scores": False,
            },
            "outputs": {
                "wetlab_shortlist_sha256": _sha256(directory / "wetlab_shortlist.csv")
            },
        }
        (directory / "manifest.json").write_text(json.dumps(manifest))

    structure_dirs = {}
    for family in ("stab", "rde"):
        directory = private / family
        directory.mkdir()
        records = []
        for peptide in EXPECTED_TARGET_SEQUENCES:
            path = directory / f"peptide_{peptide}.parquet"
            pd.DataFrame(
                {
                    "pair_uid": [f"all4-{peptide}-{rank}" for rank in range(1, 11)],
                    f"{family}_probability": [0.5] * 10,
                    f"{family}_seed_probability_sd": [0.01] * 10,
                }
            ).to_parquet(path, index=False)
            records.append(
                {"peptide_design_code": peptide, "output_sha256": _sha256(path)}
            )
        manifest = {
            "schema_version": "libb-fixed-structure-candidate-scores-merged-v1",
            "family": family,
            "candidate_universe": {"manifest_sha256": candidate_manifest_sha},
            "outputs": {"partitions": records},
            "outcome_access": {
                "selection_labels_read": False,
                "retention_measurements_read": False,
            },
        }
        (directory / "manifest.json").write_text(json.dumps(manifest))
        structure_dirs[family] = directory
    structure_manifest_hashes = {
        family: _sha256(directory / "manifest.json")
        for family, directory in structure_dirs.items()
    }

    common_mint_hash = "1" * 64
    common_mint_receipt_hash = "2" * 64
    shortlist_manifests = {
        label: json.loads(
            (private / label.replace("+", "_").lower() / "manifest.json").read_text()
        )
        for label in handoff.EXPECTED_LOCK_IDS
    }
    for manifest in shortlist_manifests.values():
        manifest["component_score_files"] = {}
    for peptide in EXPECTED_TARGET_SEQUENCES:
        stab_hash = _sha256(structure_dirs["stab"] / f"peptide_{peptide}.parquet")
        rde_hash = _sha256(structure_dirs["rde"] / f"peptide_{peptide}.parquet")
        stab_receipt_hash = structure_manifest_hashes["stab"]
        rde_receipt_hash = structure_manifest_hashes["rde"]
        shortlist_manifests["ensemble"]["component_score_files"][peptide] = {
            "additive_7site": {
                "score_sha256": "5" * 64,
                "producer_receipt_sha256": "6" * 64,
            },
            "mint_layer5": {
                "score_sha256": common_mint_hash,
                "producer_receipt_sha256": common_mint_receipt_hash,
            },
            "stab_designed_ordered": {
                "score_sha256": stab_hash,
                "producer_receipt_sha256": stab_receipt_hash,
            },
            "rde_network_designed_3fold": {
                "score_sha256": rde_hash,
                "producer_receipt_sha256": rde_receipt_hash,
            },
        }
        shortlist_manifests["MINT"]["component_score_files"][peptide] = {
            "mint_layer5": {
                "score_sha256": common_mint_hash,
                "producer_receipt_sha256": common_mint_receipt_hash,
            }
        }
        shortlist_manifests["MINT+StaB"]["component_score_files"][peptide] = {
            "mint_layer5": {
                "score_sha256": common_mint_hash,
                "producer_receipt_sha256": common_mint_receipt_hash,
            },
            "stab_designed_ordered": {
                "score_sha256": stab_hash,
                "producer_receipt_sha256": stab_receipt_hash,
            },
        }
    for label, manifest in shortlist_manifests.items():
        path = private / label.replace("+", "_").lower() / "manifest.json"
        path.write_text(json.dumps(manifest))

    return SimpleNamespace(
        ensemble_dir=private / "ensemble",
        mint_dir=private / "mint",
        mint_stab_dir=private / "mint_stab",
        stab_scores_dir=structure_dirs["stab"],
        rde_scores_dir=structure_dirs["rde"],
        candidate_universe_dir=universe,
        output_dir=private / "handoff",
        selected_peptides=None,
    )


def test_end_to_end_handoff_uses_mint_primary_and_separate_all4_alternative(
    tmp_path, monkeypatch
):
    args = _make_handoff_fixture(tmp_path, monkeypatch)
    handoff.run(args)
    output = args.output_dir
    menu = pd.read_csv(output / "all4_annotated_candidate_menu_all_12_targets.csv")
    mint_all12 = pd.read_csv(output / "primary_mint_candidate_menu_all_12_targets.csv")
    proposed = pd.read_csv(output / "primary_mint_wetlab_batch_10_targets.csv")
    order = pd.read_csv(output / "ORDER_THIS_BATCH.csv")
    result_entry = pd.read_csv(
        output / "primary_mint_result_entry_template_100_pair_assays.csv"
    )
    alternative = pd.read_csv(
        output / "primary_mint_comparator_alternative_batch_10_targets.csv"
    )
    all4_alternative = pd.read_csv(
        output / "all4_stacker_alternative_batch_10_targets.csv"
    )
    assert len(menu) == 120
    assert len(mint_all12) == 120
    assert mint_all12["pair_uid"].str.startswith("mint-").all()
    assert len(proposed) == 100
    assert proposed.groupby("peptide_design_code").size().eq(10).all()
    assert proposed["pair_uid"].str.startswith("mint-").all()
    for peptide in DEFAULT_PROPOSED_TARGETS:
        block = proposed.loc[proposed.peptide_design_code.eq(peptide)]
        assert block["pair_uid"].tolist() == [
            f"mint-{peptide}-{rank}" for rank in range(1, 11)
        ]
        assert block["model_rank"].tolist() == list(range(1, 11))
        assert block["assay_slot"].tolist() == list(range(1, 11))
    assert menu["pair_uid"].str.startswith("all4-").all()
    assert menu["candidate_menu_source_model"].eq("locked_all4_stacker").all()
    assert all4_alternative["pair_uid"].str.startswith("all4-").all()
    assert proposed["peptide_design_code"].drop_duplicates().tolist() == list(
        DEFAULT_PROPOSED_TARGETS
    )
    assert {
        "peptide_9mer_sequence",
        "model_input_smart_hla_linker_peptide_sequence",
        "provider_displayed_58aa_affibody_sequence",
        "model_input_affibody_sequence",
        "affibody_code_position_mapping",
        "mint_layer5_locked_score",
        "wetlab_rank",
        "observed_in_any_raw_round",
        "observed_in_r009_or_r010",
        "high_confidence_weak_negative",
        "affibody_identity_seen_in_strict_training",
    }.issubset(proposed.columns)
    assert not _forbidden_outcome_columns(proposed.columns)
    assert set(alternative["candidate_source"]) == {
        "primary_locked_mint_layer5_shortlist",
        "selection_supported_comparator_CMP",
    }
    assert alternative.groupby("peptide_design_code").size().eq(10).all()
    for peptide in DEFAULT_PROPOSED_TARGETS:
        expected = proposed.loc[
            proposed.peptide_design_code.eq(peptide) & proposed.wetlab_rank.lt(10), "pair_uid"
        ].tolist()
        observed = alternative.loc[
            alternative.peptide_design_code.eq(peptide)
            & alternative.assay_slot.lt(10), "pair_uid"
        ].tolist()
        assert observed == expected
        comparator = alternative.loc[
            alternative.peptide_design_code.eq(peptide)
            & alternative.assay_slot.eq(10)
        ].iloc[0]
        assert pd.isna(comparator["model_rank"])
        assert comparator["rank_semantics"] == "CMP_selection_comparator_blank_model_rank"

    assert len(order) == 100
    assert order["sample_id"].is_unique
    assert order.iloc[0]["sample_id"] == "LIBB-AF-01"
    assert order["sample_id"].str.match(r"^LIBB-[A-Z]{2}-\d{2}$").all()
    assert order["provider_displayed_58aa_affibody_sequence"].str.len().eq(58).all()
    assert all(
        assay_sequence.endswith(peptide)
        for assay_sequence, peptide in zip(
            order["model_input_smart_hla_linker_peptide_sequence"],
            order["peptide_9mer_sequence"],
        )
    )
    assert "affibody_full_sequence" not in order.columns
    assert "smart_hla_linker_peptide_full_assay_chain_sequence" not in order.columns
    assert set(order["batch_design"]) == {
        "100 specified pair assays; 16 unique displayed/model-input Affibody sequences; NOT a Cartesian cross"
    }
    assert order["sequence_readiness_warning"].str.contains(
        "NOT vendor/cloning-ready", regex=False
    ).all()
    assert len(result_entry) == 100
    assert result_entry[["sample_id", "pair_uid"]].equals(order[["sample_id", "pair_uid"]])
    result_fields = [
        "retention_at_30min_percent", "binder_ge_75", "technical_failure_status",
        "replicate_id", "assay_date", "experimental_batch_id", "notes",
    ]
    assert result_entry[result_fields].isna().all().all()
    assert len(pd.read_csv(
        output / "primary_mint_unique_affibody_sequence_review_roster.csv"
    )) == 16
    comparator_roster = pd.read_csv(
        output
        / "primary_mint_comparator_alternative_unique_affibody_sequence_review_roster.csv"
    )
    all4_roster = pd.read_csv(
        output / "all4_stacker_alternative_unique_affibody_sequence_review_roster.csv"
    )
    assert len(comparator_roster) == 10
    assert len(all4_roster) == 10
    mint_qc = pd.read_csv(output / "primary_mint_qc_summary_all_12_targets.csv")
    assert len(mint_qc) == 12
    assert {
        "previously_unobserved_candidates",
        "candidates_observed_in_r009_or_r010",
        "candidates_with_affibody_identity_seen_in_strict_training",
        "candidates_flagged_as_prior_high_confidence_weak_negative",
    }.issubset(mint_qc.columns)
    assert (output / "primary_mint_by_target_candidate_menu" / "primary_mint_peptide_AF.csv").is_file()
    assert (output / "all4_by_target_candidate_menu" / "all4_peptide_AF.csv").is_file()
    assert not (output / "candidate_menu_all_12_targets.csv").exists()
    assert not (output / "target_summary.csv").exists()
    assert not (output / "candidate_selection_audit.csv").exists()
    workbook = openpyxl.load_workbook(output / "libb_wetlab_candidate_handoff.xlsx")
    assert {
        "ORDER_THIS_BATCH", "README", "RESULT_ENTRY", "primary_mint_batch",
        "mint_comparator_alt", "all4_alt_batch", "mint_all12_menu",
        "all4_12target_menu", "mint_AF", "all4_AF",
    }.issubset(
        workbook.sheetnames
    )
    assert workbook.sheetnames[0] == "ORDER_THIS_BATCH"
    result_sheet = workbook["RESULT_ENTRY"]
    result_headers = [cell.value for cell in result_sheet[1]]
    for field in result_fields:
        column_index = result_headers.index(field) + 1
        assert all(
            result_sheet.cell(row=row_index, column=column_index).value is None
            for row_index in range(2, 102)
        )
    readme = list(workbook["README"].values)
    readme_text = " ".join(str(value) for row in readme for value in row if value is not None)
    assert "100 specified peptide-Affibody pair assays" in readme_text
    assert "NOT a Cartesian cross" in readme_text
    assert "NOT vendor- or cloning-ready" in readme_text
    assert "do NOT sort" in readme_text
    assert "pair identities were used upstream only to exclude" in readme_text
    manifest = json.loads((output / "manifest.json").read_text())
    assert manifest["schema_version"] == "libb-wetlab-candidate-handoff-v4"
    assert manifest["proposed_batch"]["mint_layer5_is_primary"] is True
    assert manifest["proposed_batch"]["all4_stacker_is_separate_alternative"] is True
    assert manifest["proposed_batch"]["pair_assays"] == 100
    assert manifest["proposed_batch"]["unique_affibody_sequences_primary_mint"] == 16
    assert manifest["proposed_batch"]["unique_affibody_sequences_comparator_alternative"] == len(
        comparator_roster
    )
    assert manifest["proposed_batch"]["unique_affibody_sequences_all4_alternative"] == len(
        all4_roster
    )
    assert "NOT a Cartesian cross" in manifest["proposed_batch"]["batch_geometry"]
    assert "not a final vendor or cloning construct" in manifest["sequence_readiness"]["affibody"]
    assert manifest["outcome_access"][
        "retention_panel_pair_identities_used_upstream_for_exclusion"
    ] is True
    assert manifest["outcome_access"][
        "retention_outcomes_used_for_model_training_scoring_ranking_cutoff_or_target_choice"
    ] is False
    assert manifest["prospective_result_definition"]["result_entry_template_released_blank"] is True
    assert manifest["validation"]["result_entry_outcome_fields_blank_at_release"] is True
    assert manifest["validation"]["no_cross_peptide_score_sorting"] is True
    assert manifest["outputs"]["workbook"]["sha256"] == _sha256(
        output / "libb_wetlab_candidate_handoff.xlsx"
    )
    assert not list(output.parent.glob(f".{output.name}.staging-*"))


def test_handoff_rejects_cross_shortlist_lock_bundle_mismatch(tmp_path, monkeypatch):
    args = _make_handoff_fixture(tmp_path, monkeypatch)
    manifest_path = args.mint_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    manifest["ensemble_lock"]["bundle_manifest_sha256"] = "9" * 64
    manifest_path.write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match="same weak-OOF lock bundle"):
        handoff.run(args)
    assert not args.output_dir.exists()


def test_handoff_rejects_shared_mint_score_hash_mismatch(tmp_path, monkeypatch):
    args = _make_handoff_fixture(tmp_path, monkeypatch)
    manifest_path = args.mint_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    manifest["component_score_files"]["AF"]["mint_layer5"]["score_sha256"] = "9" * 64
    manifest_path.write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match="shared mint_layer5 score hashes differ"):
        handoff.run(args)
    assert not args.output_dir.exists()


def test_handoff_rejects_tampered_excluded_pairs(tmp_path, monkeypatch):
    args = _make_handoff_fixture(tmp_path, monkeypatch)
    excluded_path = args.candidate_universe_dir / "excluded_pairs.parquet"
    frame = pd.read_parquet(excluded_path)
    frame.loc[0, "pooled_r009_r010_count"] += 1
    frame.to_parquet(excluded_path, index=False)
    with pytest.raises(ValueError, match="hash differs"):
        handoff.run(args)
    assert not args.output_dir.exists()


def test_handoff_rejects_hidden_retention_column_in_structure_scores(
    tmp_path, monkeypatch
):
    args = _make_handoff_fixture(tmp_path, monkeypatch)
    score_path = args.stab_scores_dir / "peptide_AF.parquet"
    scores = pd.read_parquet(score_path)
    scores["observed_retention_percent"] = 99.0
    scores.to_parquet(score_path, index=False)
    score_hash = _sha256(score_path)

    structure_manifest_path = args.stab_scores_dir / "manifest.json"
    structure_manifest = json.loads(structure_manifest_path.read_text())
    for record in structure_manifest["outputs"]["partitions"]:
        if record["peptide_design_code"] == "AF":
            record["output_sha256"] = score_hash
    structure_manifest_path.write_text(json.dumps(structure_manifest))
    receipt_hash = _sha256(structure_manifest_path)

    for directory in (args.ensemble_dir, args.mint_stab_dir):
        manifest_path = directory / "manifest.json"
        manifest = json.loads(manifest_path.read_text())
        for peptide, sources in manifest["component_score_files"].items():
            sources["stab_designed_ordered"]["producer_receipt_sha256"] = receipt_hash
            if peptide == "AF":
                sources["stab_designed_ordered"]["score_sha256"] = score_hash
        manifest_path.write_text(json.dumps(manifest))

    with pytest.raises(ValueError, match="leaks retention outcomes"):
        handoff.run(args)
    assert not args.output_dir.exists()


def test_handoff_does_not_publish_partial_output_if_workbook_fails(
    tmp_path, monkeypatch
):
    args = _make_handoff_fixture(tmp_path, monkeypatch)

    def fail_workbook(*_args, **_kwargs):
        raise RuntimeError("injected workbook failure")

    monkeypatch.setattr(handoff, "write_wetlab_workbook", fail_workbook)
    with pytest.raises(RuntimeError, match="injected workbook failure"):
        handoff.run(args)
    assert not args.output_dir.exists()
